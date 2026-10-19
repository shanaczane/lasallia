# Reads the LRC's Excel catalog sheets (one sheet per program, one row per
# copy) into the same records the MARC importer produces.
#
# A row with only an accession number is another copy of the book above it.
# Errors block a row; warnings flag likely typos but still import it.

import io
import re

import openpyxl

from core.accession import normalize_accession_number
from core.importers.marc import ImportIssue, ParsedRecord, normalize_isbn
from scripts.excel_source import COLUMN_ALIASES

TEMPLATE_COLUMNS = list(COLUMN_ALIASES.keys())
NOTES_SHEET = "How to fill this in"  # build_template's instructions sheet


def _text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # ISBNs/years typed into a number-formatted cell
    cleaned = re.sub(r"\s+", " ", str(value)).strip()
    return cleaned or None


def _publisher(place_of_publication: str | None) -> str | None:
    """'Quezon City : Phoenix Publishing House.' -> 'Phoenix Publishing House'
    (same rule as scripts/seed_books.py's split_publisher)."""
    if not place_of_publication:
        return None
    _, _, publisher = place_of_publication.partition(":")
    publisher = (publisher or place_of_publication).strip().rstrip(".,").strip()
    return publisher or None


def _group_key(title: str, author: str | None, isbn: str | None) -> str:
    return re.sub(r"[^a-z0-9]", "", f"{title}|{author or ''}|{normalize_isbn(isbn) or ''}".lower())


def parse_spreadsheet(data: bytes, college: str | None) -> tuple[list[ParsedRecord], list[ImportIssue], list[ImportIssue]]:
    """Returns (records, errors, warnings). Rows for the same title (same
    title + author + ISBN) become one record with several copies — the
    sheets list each physical copy on its own row."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ValueError("This isn't a readable .xlsx file. Save it from Excel as an .xlsx workbook and try again.")

    errors: list[ImportIssue] = []
    warnings: list[ImportIssue] = []
    records: dict[str, ParsedRecord] = {}
    seen_accessions: dict[str, str] = {}
    row_count = 0

    for sheet in wb.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        header = [_text(h) for h in header]
        col: dict[str, int] = {}
        for canonical, aliases in COLUMN_ALIASES.items():
            for alias in aliases:
                if alias in header:
                    col[canonical] = header.index(alias)
                    break
        if "Title" not in col or "Book ID (Accession No.)" not in col:
            if sheet.title == NOTES_SHEET:
                continue
            warnings.append(ImportIssue(0, None, "Sheet skipped — no \"Title\" and \"Book ID (Accession No.)\" header row", f"Sheet {sheet.title}"))
            continue

        previous: ParsedRecord | None = None  # last title row on this sheet
        for row_no, row in enumerate(rows, start=2):
            cell = {name: _text(row[i]) if i < len(row) else None for name, i in col.items()}
            if not any(cell.values()):
                continue  # blank spacer row
            row_count += 1
            where = f"Sheet {sheet.title}, row {row_no}"
            title = cell.get("Title")
            accession = normalize_accession_number(cell.get("Book ID (Accession No.)") or "")

            if not accession:
                errors.append(ImportIssue(row_count, title, "No accession number (Book ID)", where))
                continue
            if accession in seen_accessions:
                errors.append(ImportIssue(row_count, title, f"Accession {accession} is already used at {seen_accessions[accession]}", where))
                continue

            # A row with only an accession number is another copy of the
            # book above it.
            if not title:
                only_accession = all(v is None for k, v in cell.items() if k != "Book ID (Accession No.)")
                if only_accession and previous is not None:
                    seen_accessions[accession] = where
                    previous.barcodes.append(accession)
                else:
                    errors.append(ImportIssue(row_count, None, f"Accession {accession} has no title", where))
                continue
            seen_accessions[accession] = where

            author = cell.get("Author(s) Full Name") or cell.get("Author(s)")
            isbn = cell.get("ISBN")
            year_text = cell.get("Year")
            year_match = re.search(r"(?<!\d)(1[5-9]\d\d|20\d\d)(?!\d)", year_text or "")

            if year_text and not year_match:
                warnings.append(ImportIssue(row_count, title, f"Year \"{year_text}\" isn't a year — left blank", where))
            if isbn and len(normalize_isbn(isbn) or "") not in (10, 13):
                warnings.append(ImportIssue(row_count, title, f"ISBN \"{isbn}\" should have 10 or 13 digits", where))
            if not cell.get("Call No."):
                warnings.append(ImportIssue(row_count, title, "No call number — shelf will show as Unassigned", where))
            if not author:
                warnings.append(ImportIssue(row_count, title, "No author — will show as Unknown", where))

            key = _group_key(title, author, isbn)
            record = records.get(key)
            if record is None:
                record = records[key] = ParsedRecord(
                    index=row_count,
                    location=where,
                    title=title,
                    author=author,
                    isbn=isbn,
                    call_number=cell.get("Call No."),
                    publisher=_publisher(cell.get("Place of Publication")),
                    published_year=int(year_match.group(1)) if year_match else None,
                    abstract=cell.get("Description"),
                    college=college,
                    program=sheet.title,
                    cover_url=cell.get("Images"),
                )
            record.barcodes.append(accession)
            previous = record

    wb.close()
    return list(records.values()), errors, warnings


def build_template() -> bytes:
    """Blank workbook in the layout parse_spreadsheet reads."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Program name"
    ws.append(TEMPLATE_COLUMNS)
    ws.append([
        "Martin, R.", "Robert C. Martin", 2008, "Clean code : a handbook of agile software craftsmanship",
        "Upper Saddle River, NJ : Prentice Hall", "005.1 M379 2008", "T00001", "Optional summary of the book.",
        "9780132350884", "",
    ])
    for column, width in zip("ABCDEFGHIJ", (16, 22, 8, 40, 30, 18, 22, 30, 16, 20)):
        ws.column_dimensions[column].width = width

    notes = wb.create_sheet(NOTES_SHEET)
    for line in (
        ["One sheet per program — the sheet's name becomes the book's program (e.g. rename \"Program name\" to \"BSIT\")."],
        ["One row per physical copy. For more copies of a book, add rows right under it with only the Book ID (Accession No.) filled in."],
        ["Required: Title and Book ID (Accession No.). Everything else is optional but recommended."],
        ["Delete the example row before uploading. This sheet can stay — it's ignored on upload."],
    ):
        notes.append(line)
    notes.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
