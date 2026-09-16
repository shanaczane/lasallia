# apps/api/scripts/import_rfid_enrollment.py
# One-time (and re-run-as-needed) import: reads the librarian-filled
# RFID/NFC card enrollment spreadsheet and saves each row's Card UID,
# Program, and Year Level onto the matching student's profile, via the
# service-role client (bypasses RLS — this is an admin operation, not a
# user-facing write).
#
# Program/Year Level were sitting right there in this same spreadsheet
# (a librarian already fills them in alongside Card UID) but were never
# read — every real student account had them null, with no way to set
# them, until the Patrons screen grew an edit form for it. This is the
# faster path for a whole roster at once instead of one-by-one there.
#
# Matches rows to students by email (profiles has no student-number
# column — email is the only reliably-unique identifying field a
# spreadsheet row can carry). Skips blank template rows, refuses to
# silently overwrite a Card UID already assigned to a DIFFERENT student (a
# unique constraint on profiles.rfid_uid would reject that anyway, but
# checking first gives a readable error instead of a raw DB exception),
# and no-ops a row that already matches what's stored in every column.
#
# Usage: venv/Scripts/python.exe scripts/import_rfid_enrollment.py [--dry-run]

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # so `core.*` resolves

import openpyxl

from core.supabase import get_admin_client

SOURCE_PATH = Path(__file__).resolve().parent.parent / "data" / "rfid_card_enrollment.xlsx"
SHEET_NAME = "Card Enrollment"


def _parse_year_level(value) -> int | None:
    """'1st Year' / '2nd Year' / ... -> 1 / 2 / ... Also accepts a bare
    number, in case a future sheet just puts '1' in the cell."""
    if value is None:
        return None
    match = re.search(r"\d+", str(value))
    return int(match.group()) if match else None


def read_rows(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: i for i, name in enumerate(header)}

    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        email = (raw[col["Email"]] or "").strip() if raw[col["Email"]] else ""
        uid = (raw[col["Card UID"]] or "").strip() if raw[col["Card UID"]] else ""
        if not email or not uid:
            continue  # blank template row — nothing to import

        program = None
        if "Program" in col:
            program = (raw[col["Program"]] or "").strip() or None
        year_level = _parse_year_level(raw[col["Year Level"]]) if "Year Level" in col else None

        rows.append({
            "full_name": raw[col["Full Name"]],
            "email": email,
            "card_uid": uid,
            "program": program,
            "year_level": year_level,
        })
    return rows


def run(dry_run: bool) -> None:
    admin = get_admin_client()
    rows = read_rows(SOURCE_PATH)

    updated, unchanged, errors = [], [], []

    for row in rows:
        profile_res = (
            admin.table("profiles")
            .select("id, full_name, rfid_uid, program, year_level")
            .eq("email", row["email"])
            .execute()
        )
        if not profile_res.data:
            errors.append(f"{row['email']}: no student account with this email exists")
            continue
        profile = profile_res.data[0]

        payload: dict = {}
        changed_fields: list[str] = []

        if row["card_uid"] and profile["rfid_uid"] != row["card_uid"]:
            conflict_res = admin.table("profiles").select("id, full_name").eq("rfid_uid", row["card_uid"]).execute()
            if conflict_res.data and conflict_res.data[0]["id"] != profile["id"]:
                other = conflict_res.data[0]["full_name"]
                errors.append(f"{row['email']}: Card UID {row['card_uid']!r} is already assigned to {other!r}")
                continue
            payload["rfid_uid"] = row["card_uid"]
            changed_fields.append(f"Card UID={row['card_uid']}")

        if row["program"] and profile["program"] != row["program"]:
            payload["program"] = row["program"]
            changed_fields.append(f"Program={row['program']}")

        if row["year_level"] and profile["year_level"] != row["year_level"]:
            payload["year_level"] = row["year_level"]
            changed_fields.append(f"Year Level={row['year_level']}")

        if not payload:
            unchanged.append(row["email"])
            continue

        if not dry_run:
            admin.table("profiles").update(payload).eq("id", profile["id"]).execute()
        updated.append(f"{row['email']}: {', '.join(changed_fields)}")

    label = "Would update" if dry_run else "Updated"
    print(f"{label} {len(updated)} student(s):")
    for line in updated:
        print(f"  - {line}")
    if unchanged:
        print(f"\nAlready up to date ({len(unchanged)}): {', '.join(unchanged)}")
    if errors:
        print(f"\n{len(errors)} row(s) need attention:")
        for e in errors:
            print(f"  ! {e}")


if __name__ == "__main__":
    run(dry_run="--dry-run" in sys.argv)
